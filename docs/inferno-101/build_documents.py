import html
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
CONTENT = json.loads((ROOT / "content.json").read_text(encoding="utf-8"))
OUTPUT = ROOT / "output"
OUTPUT.mkdir(exist_ok=True)

ORANGE = "FF6A21"
DARK = "10283B"
PALE = "FFF8EC"
LIGHT = "F3F5F6"
LINE = "D9D9D9"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_borders(cell, color=LINE, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def mark_row_as_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def style_run(run, size=None, bold=None, color=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_labelled_paragraph(doc, label, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    label_run = paragraph.add_run(label + "  ")
    style_run(label_run, size=9, bold=True, color=ORANGE)
    text_run = paragraph.add_run(text)
    style_run(text_run, size=10.5, color="222222")
    return paragraph


def add_bullets(doc, items, label=None):
    if label:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(label)
        style_run(r, size=9, bold=True, color=ORANGE)
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        style_run(r, size=10.2, color="222222")


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(34, 34, 34)
    for style_name, size in (("Title", 28), ("Heading 1", 18), ("Heading 2", 13)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    logo = ROOT / "assets" / "sparked-logo.png"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    logo_shape = p.add_run().add_picture(str(logo), width=Inches(2.05))
    logo_shape._inline.docPr.set("descr", "Sparked HL7 FHIR logo")

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(28)
    title.paragraph_format.space_after = Pt(8)
    title.add_run(CONTENT["title"] + " Recording Script")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run(CONTENT["subtitle"])
    style_run(r, size=15, color=ORANGE)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(12)
    r = intro.add_run(
        "A screen-led draft for team review. Follow the screen directions while reading the spoken copy naturally. "
        "The target runtime is fourteen minutes, with extra detail available if questions are recorded separately."
    )
    style_run(r, size=11)

    meta = doc.add_table(rows=2, cols=3)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    widths = [2.0, 2.8, 2.0]
    values = [
        ("Presenter", CONTENT["presenter"]),
        ("Audience", CONTENT["audience"]),
        ("Environment", CONTENT["environment"])
    ]
    mark_row_as_header(meta.rows[0])
    for index, (label, value) in enumerate(values):
        for row_index, text in enumerate((label, value)):
            cell = meta.cell(row_index, index)
            cell.width = Inches(widths[index])
            set_cell_fill(cell, PALE if row_index == 0 else LIGHT)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(text)
            style_run(run, size=8.5 if row_index == 0 else 9.5, bold=(row_index == 0), color=ORANGE if row_index == 0 else "222222")

    production_line = doc.add_paragraph()
    production_line.paragraph_format.space_before = Pt(7)
    production_line.paragraph_format.space_after = Pt(1)
    for label, value in (
        ("Target", CONTENT["target_length"]),
        ("Format", "Slides and live browser demo"),
        ("Data", "Synthetic only"),
    ):
        label_run = production_line.add_run(label + ": ")
        style_run(label_run, size=9, bold=True, color=ORANGE)
        value_run = production_line.add_run(value + "   ")
        style_run(value_run, size=9, color="222222")

    doc.add_paragraph()
    heading = doc.add_paragraph("Run order", style="Heading 1")
    heading.paragraph_format.space_before = Pt(4)
    heading.paragraph_format.space_after = Pt(8)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    mark_row_as_header(table.rows[0])
    hdr = table.rows[0].cells
    for index, text in enumerate(("Segment", "Screen", "Time")):
        hdr[index].width = Inches((2.0, 4.2, 0.9)[index])
        set_cell_fill(hdr[index], DARK)
        set_cell_borders(hdr[index])
        set_cell_margins(hdr[index])
        run = hdr[index].paragraphs[0].add_run(text)
        style_run(run, size=9, bold=True, color="FFFFFF")
    for idx, segment in enumerate(CONTENT["segments"]):
        row = table.add_row().cells
        values = [f'{segment["number"]}. {segment["title"]}', segment["screen"], segment["time"]]
        for col, value in enumerate(values):
            row[col].width = Inches((2.0, 4.2, 0.9)[col])
            set_cell_fill(row[col], "FFFFFF" if idx % 2 == 0 else PALE)
            set_cell_borders(row[col])
            set_cell_margins(row[col])
            row[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = row[col].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            style_run(run, size=9.2, bold=(col == 0), color="222222")

    for segment in CONTENT["segments"]:
        heading = doc.add_paragraph(
            f'{segment["number"]}  {segment["title"]}  {segment["duration"]}',
            style="Heading 1"
        )
        heading.paragraph_format.page_break_before = segment["number"] in (5, 6, 7)
        heading.paragraph_format.space_before = Pt(10)
        heading.paragraph_format.space_after = Pt(7)
        add_labelled_paragraph(doc, "SCREEN", segment["screen"])
        if segment.get("prep"):
            add_bullets(doc, segment["prep"], "PREP")
        if segment.get("actions"):
            add_bullets(doc, segment["actions"], "ACTIONS")
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("SAY")
        style_run(r, size=9, bold=True, color=ORANGE)
        for paragraph_text in segment["say"]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.right_indent = Inches(0.1)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(paragraph_text)
            style_run(r, size=10.6, color="111111")
        if segment.get("fallback"):
            add_labelled_paragraph(doc, "FALLBACK", segment["fallback"])

    doc.add_page_break()
    doc.add_paragraph("Recording day preflight", style="Heading 1")
    add_bullets(doc, CONTENT["preflight"])
    doc.add_paragraph("Prepared rehearsal sessions", style="Heading 2")
    for label, url in CONTENT["rehearsal_sessions"]:
        add_labelled_paragraph(doc, label.upper(), url)

    doc.core_properties.title = CONTENT["title"] + " Recording Script"
    doc.core_properties.subject = "Draft Sparked FHIR AU Inferno Test Kit 101 recording"
    doc.core_properties.author = "Sparked AU FHIR Accelerator"
    out = OUTPUT / "Inferno 101 Draft Script.docx"
    doc.save(out)
    return out


def render_list(items, class_name=""):
    cls = f' class="{class_name}"' if class_name else ""
    return "<ul" + cls + ">" + "".join(f"<li>{html.escape(item)}</li>" for item in items) + "</ul>"


def build_html():
    rows = "".join(
        f'<tr><td><a href="#segment-{s["number"]}">{s["number"]}. {html.escape(s["title"])}</a></td>'
        f'<td>{html.escape(s["screen"])}</td><td>{html.escape(s["time"])}</td></tr>'
        for s in CONTENT["segments"]
    )
    sections = []
    for s in CONTENT["segments"]:
        prep = f'<div class="cue"><h3>Prep</h3>{render_list(s["prep"])}</div>' if s.get("prep") else ""
        actions = f'<div class="cue"><h3>Actions</h3>{render_list(s["actions"])}</div>' if s.get("actions") else ""
        spoken = "".join(f"<p>{html.escape(p)}</p>" for p in s["say"])
        fallback = (
            f'<div class="fallback"><strong>Fallback</strong><p>{html.escape(s["fallback"])}</p></div>'
            if s.get("fallback") else ""
        )
        sections.append(f'''
        <section class="segment" id="segment-{s["number"]}" data-start="{html.escape(s["time"].split(" to ")[0])}">
          <header class="segment-head">
            <div><span class="segment-number">{s["number"]}</span><h2>{html.escape(s["title"])}</h2></div>
            <span class="duration">{html.escape(s["duration"])}</span>
          </header>
          <div class="screen"><strong>Screen</strong><p>{html.escape(s["screen"])}</p></div>
          {prep}
          {actions}
          <div class="say"><h3>Say</h3>{spoken}</div>
          {fallback}
          <label class="done"><input type="checkbox"> Segment rehearsed</label>
        </section>''')
    session_rows = "".join(
        f'<tr><td>{html.escape(label)}</td><td><a href="{html.escape(url)}">{html.escape(url)}</a></td></tr>'
        for label, url in CONTENT["rehearsal_sessions"]
    )
    preflight = "".join(
        f'<li><label><input type="checkbox"> {html.escape(item)}</label></li>' for item in CONTENT["preflight"]
    )
    template = '''<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>__TITLE__ Run Sheet</title>
<style>
:root { --ink:#14212b; --muted:#5c6670; --cream:#fff9ec; --orange:#ff6a21; --navy:#10283b; --line:#d9dde0; --paper:#fff; --green:#277a55; }
* { box-sizing:border-box; }
html { scroll-behavior:smooth; }
body { margin:0; color:var(--ink); background:#f4f2eb; font:16px/1.5 Arial, Helvetica, sans-serif; }
a { color:#005f86; }
.hero { background:var(--cream); border-top:10px solid var(--orange); padding:34px max(24px, calc((100vw - 1080px)/2)); }
.brand { width:220px; height:auto; display:block; margin-bottom:28px; }
.eyebrow { color:#a54116; font-size:.82rem; font-weight:700; letter-spacing:.08em; text-transform:uppercase; }
h1 { margin:.25rem 0 .4rem; font-size:clamp(2.3rem, 6vw, 4.6rem); line-height:1; letter-spacing:-.035em; }
.lede { margin:0; max-width:760px; color:var(--muted); font-size:1.2rem; }
.meta { display:flex; flex-wrap:wrap; gap:18px 34px; margin-top:24px; font-size:.92rem; }
.meta strong { display:block; color:var(--orange); font-size:.72rem; letter-spacing:.07em; text-transform:uppercase; }
.timerbar { position:sticky; top:0; z-index:20; display:flex; align-items:center; gap:12px; padding:10px max(18px, calc((100vw - 1080px)/2)); color:#fff; background:var(--navy); box-shadow:0 2px 12px #0002; }
.timer { min-width:86px; font:700 1.35rem/1 monospace; }
button { border:1px solid #ffffff66; border-radius:5px; padding:7px 12px; color:#fff; background:transparent; font-weight:700; cursor:pointer; }
button.primary { border-color:var(--orange); background:var(--orange); }
.timer-note { margin-left:auto; color:#dbe6ed; font-size:.85rem; }
main { width:min(1080px, calc(100% - 32px)); margin:28px auto 80px; }
.overview, .preflight, .sessions, .segment { margin:0 0 24px; border:1px solid var(--line); border-radius:9px; background:var(--paper); box-shadow:0 5px 18px #10283b0c; overflow:hidden; }
.overview, .preflight, .sessions { padding:24px; }
h2 { margin:0; font-size:1.55rem; line-height:1.15; }
h3 { margin:0 0 8px; color:#9d3e13; font-size:.78rem; letter-spacing:.08em; text-transform:uppercase; }
table { width:100%; border-collapse:collapse; font-size:.93rem; }
th { color:#fff; background:var(--navy); text-align:left; }
th, td { border:1px solid var(--line); padding:10px 12px; vertical-align:top; }
th:last-child, td:last-child { width:112px; white-space:nowrap; }
.checklist { margin:0; padding:0; list-style:none; columns:2; column-gap:36px; }
.checklist li { break-inside:avoid; margin:0 0 9px; }
input[type=checkbox] { width:17px; height:17px; vertical-align:-3px; accent-color:var(--orange); }
.segment-head { display:flex; justify-content:space-between; align-items:center; padding:18px 22px; color:#fff; background:var(--navy); }
.segment-head > div { display:flex; align-items:center; gap:13px; }
.segment-number { display:grid; place-items:center; width:34px; height:34px; border-radius:50%; color:var(--navy); background:var(--orange); font-weight:800; }
.duration { font:700 1rem monospace; }
.screen, .cue, .say, .fallback { padding:18px 22px; border-bottom:1px solid var(--line); }
.screen { display:grid; grid-template-columns:90px 1fr; gap:12px; background:var(--cream); }
.screen strong { color:#a54116; font-size:.78rem; letter-spacing:.08em; text-transform:uppercase; }
.screen p, .fallback p { margin:0; }
.cue ul { margin:0; padding-left:20px; }
.cue li { margin:.25rem 0; }
.say { border-left:7px solid var(--orange); }
.say p { max-width:790px; margin:.5rem 0 1rem; font-size:1.05rem; }
.fallback { background:#f4f6f7; }
.fallback strong { color:#7a350f; }
.done { display:block; padding:13px 22px; color:var(--green); font-weight:700; background:#f5fbf7; }
.sessions td:first-child { width:220px; font-weight:700; }
.sessions a { overflow-wrap:anywhere; }
.footer { width:min(1080px, calc(100% - 32px)); margin:0 auto 50px; color:var(--muted); font-size:.86rem; }
@media (max-width:760px) { .checklist { columns:1; } .timer-note { display:none; } .screen { grid-template-columns:1fr; } th:nth-child(2), td:nth-child(2) { display:none; } }
@media print { body { background:#fff; font-size:10pt; } .hero { padding:18px 0; } .brand { width:150px; margin-bottom:12px; } .timerbar, .done { display:none; } main { width:100%; margin:12px 0; } .overview, .preflight, .sessions, .segment { box-shadow:none; break-inside:avoid; } .segment { break-before:page; } a { color:#000; text-decoration:none; } }
</style>
</head>
<body>
<header class="hero">
  <img class="brand" src="../assets/sparked-logo.png" alt="Sparked HL7 FHIR">
  <div class="eyebrow">Sparked FHIR AU recording draft</div>
  <h1>__TITLE__</h1>
  <p class="lede">__SUBTITLE__. A complete cue sheet for a draft recording and team review.</p>
  <div class="meta">
    <span><strong>Presenter</strong>__PRESENTER__</span>
    <span><strong>Audience</strong>__AUDIENCE__</span>
    <span><strong>Target</strong>__TARGET__</span>
    <span><strong>Environment</strong><a href="__ENV__">inferno.hl7.org.au</a></span>
  </div>
</header>
<div class="timerbar">
  <span class="timer" id="timer">00:00</span>
  <button class="primary" id="start">Start</button>
  <button id="reset">Reset</button>
  <span class="timer-note">Space starts or pauses. R resets.</span>
</div>
<main>
  <section class="overview">
    <h2>Run order</h2>
    <p>Slides establish the context. The browser demonstration carries the middle of the recording.</p>
    <table><thead><tr><th>Segment</th><th>Screen</th><th>Time</th></tr></thead><tbody>__ROWS__</tbody></table>
  </section>
  <section class="preflight">
    <h2>Recording day preflight</h2>
    <p>Complete this immediately before recording. These checks contain the details most likely to change.</p>
    <ul class="checklist">__PREFLIGHT__</ul>
  </section>
  __SECTIONS__
  <section class="sessions">
    <h2>Prepared rehearsal sessions</h2>
    <p>Use these only as fallbacks. Public sessions can be purged, and live results can change.</p>
    <table><thead><tr><th>Run</th><th>Session URL</th></tr></thead><tbody>__SESSION_ROWS__</tbody></table>
  </section>
</main>
<p class="footer">Draft for team review. Source content verified against the planned Inferno guidance copy on 11 September 2026.</p>
<script>
let elapsed = 0;
let startedAt = null;
let timerId = null;
const timer = document.getElementById('timer');
const start = document.getElementById('start');
function paint() {
  const total = elapsed + (startedAt ? Date.now() - startedAt : 0);
  const seconds = Math.floor(total / 1000);
  timer.textContent = String(Math.floor(seconds / 60)).padStart(2, '0') + ':' + String(seconds % 60).padStart(2, '0');
}
function toggle() {
  if (startedAt) {
    elapsed += Date.now() - startedAt;
    startedAt = null;
    clearInterval(timerId);
    timerId = null;
    start.textContent = 'Resume';
  } else {
    startedAt = Date.now();
    timerId = setInterval(paint, 250);
    start.textContent = 'Pause';
  }
  paint();
}
function resetTimer() {
  elapsed = 0;
  startedAt = null;
  clearInterval(timerId);
  timerId = null;
  start.textContent = 'Start';
  paint();
}
start.addEventListener('click', toggle);
document.getElementById('reset').addEventListener('click', resetTimer);
document.addEventListener('keydown', event => {
  if (event.target.matches('input, button, a')) return;
  if (event.code === 'Space') { event.preventDefault(); toggle(); }
  if (event.key.toLowerCase() === 'r') resetTimer();
});
</script>
</body>
</html>'''
    replacements = {
        "__TITLE__": html.escape(CONTENT["title"]),
        "__SUBTITLE__": html.escape(CONTENT["subtitle"]),
        "__PRESENTER__": html.escape(CONTENT["presenter"]),
        "__AUDIENCE__": html.escape(CONTENT["audience"]),
        "__TARGET__": html.escape(CONTENT["target_length"]),
        "__ENV__": html.escape(CONTENT["environment"]),
        "__ROWS__": rows,
        "__PREFLIGHT__": preflight,
        "__SECTIONS__": "".join(sections),
        "__SESSION_ROWS__": session_rows
    }
    for key, value in replacements.items():
        template = template.replace(key, value)
    template = "\n".join(line.rstrip() for line in template.splitlines()) + "\n"
    out = OUTPUT / "Inferno 101 Run Sheet.html"
    out.write_text(template, encoding="utf-8")
    return out


if __name__ == "__main__":
    print(build_docx())
    print(build_html())
